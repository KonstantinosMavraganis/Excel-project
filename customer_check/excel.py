import re
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .models import Result, Row, Table


def _open(path: str):
    if Path(path).suffix.lower() not in {".xlsx", ".xlsm"}:
        raise ValueError("Υποστηρίζονται .xlsx και .xlsm. Αποθήκευσε παλιό .xls ως .xlsx στο Excel.")
    return load_workbook(path, read_only=True, data_only=True)


def sheets(path: str) -> list[str]:
    workbook = _open(path)
    try:
        return workbook.sheetnames
    finally:
        workbook.close()


def _text(cell) -> str:
    value = cell.value
    if value is None:
        return ""
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        if isinstance(value, int) or value.is_integer():
            value = int(value)
            # Excel often stores a code as a number with a display format like 000000.
            if re.fullmatch(r"0+", cell.number_format or ""):
                return str(value).zfill(len(cell.number_format))
            return str(value)
    return str(value).strip()


def read_table(path: str, sheet: str, header_row: int = 1) -> Table:
    if header_row < 1:
        raise ValueError("Η γραμμή επικεφαλίδων πρέπει να είναι τουλάχιστον 1.")
    workbook = _open(path)
    try:
        worksheet = workbook[sheet]
        iterator = worksheet.iter_rows(min_row=header_row)
        header_cells = next(iterator, ())
        if not header_cells or not any(_text(cell) for cell in header_cells):
            raise ValueError("Δεν βρέθηκαν επικεφαλίδες στην επιλεγμένη γραμμή.")
        headers = tuple(f"{i + 1}: {_text(cell) or '(χωρίς τίτλο)'}" for i, cell in enumerate(header_cells))
        rows = []
        for number, cells in enumerate(iterator, header_row + 1):
            values = tuple(_text(cell) for cell in cells)
            if any(values):
                rows.append(Row(number, values))
        return Table(str(Path(path).resolve()), sheet, headers, tuple(rows))
    finally:
        workbook.close()


def export_report(path: str, result: Result, metadata: dict[str, str]) -> None:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Σύνοψη"
    summary.append(["Πεδίο", "Τιμή"])
    for key, value in metadata.items():
        summary.append([key, value])
    for label, value in [
        ("Ενεργοί με εγγραφή πληρωμής", len(result.found)),
        ("Ενεργοί που λείπουν", len(result.missing)),
        ("Άγνωστοι κωδικοί πληρωμών", len(result.unknown)),
        ("Εγγραφές προς έλεγχο", len(result.issues)),
        ("Γραμμές μητρώου", result.registry_rows),
        ("Γραμμές πληρωμών", result.payment_rows),
        ("Μη ενεργές γραμμές", result.inactive_rows),
        ("Επιπλέον εγγραφές ίδιου κωδικού στις πληρωμές", result.duplicate_payments),
    ]:
        summary.append([label, value])
    summary.append(["Σημείωση", "Ελέγχεται η παρουσία κωδικού, όχι το ποσό ή η πλήρης εξόφληση."])
    summary.append(["Τύποι Excel", "Χρησιμοποιούνται αποθηκευμένα αποτελέσματα τύπων. Ανανέωσε και αποθήκευσε τα αρχεία στο Excel πριν τον έλεγχο."])
    for title, entries in [
        ("Λείπουν", result.missing), ("Βρέθηκαν", result.found),
        ("Άγνωστοι κωδικοί", result.unknown), ("Προς έλεγχο", result.issues),
    ]:
        sheet = workbook.create_sheet(title)
        sheet.append(["Κωδικός πελάτη", "Όνομα / Επωνυμία", "Γραμμή προέλευσης", "Λεπτομέρεια"])
        for entry in entries:
            sheet.append([entry.code, entry.name, entry.source_row, entry.detail])
        sheet.auto_filter.ref = sheet.dimensions
    for sheet in workbook:
        sheet.freeze_panes = "A2"
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="234E70")
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                # Uploaded strings must remain text, even if they start with '='.
                if isinstance(cell.value, str):
                    cell.data_type = "s"
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        for column in sheet.columns:
            letter = column[0].column_letter
            sheet.column_dimensions[letter].width = min(65, max(18, max(len(str(c.value or "")) for c in column) + 2))
    workbook.save(path)
    workbook.close()

