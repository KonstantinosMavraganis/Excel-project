from pathlib import Path
from tempfile import TemporaryDirectory
import unittest

from openpyxl import Workbook, load_workbook

from customer_check.excel import export_report, read_table, sheets
from customer_check.models import Entry, Result, Row, Table
from customer_check.reconcile import compare


def table(rows, headers=("Κωδικός", "Συμβόλαιο", "Όνομα")):
    return Table("example.xlsx", "Sheet", headers, tuple(Row(i + 2, tuple(row)) for i, row in enumerate(rows)))


class ReconciliationTests(unittest.TestCase):
    def check(self, registry, payments):
        return compare(table(registry), table([(code,) for code in payments], ("Κωδικός",)), 0, 1, 0, {"Ναι"}, 2)

    def test_active_missing_and_unknown(self):
        result = self.check([("001", "Ναι", "Α"), ("002", "Ναι", "Β"), ("003", "Όχι", "Γ")], ["001", "999"])
        self.assertEqual([e.code for e in result.found], ["001"])
        self.assertEqual([e.code for e in result.missing], ["002"])
        self.assertEqual([e.code for e in result.unknown], ["999"])
        self.assertEqual(result.inactive_rows, 1)

    def test_whitespace_case_and_duplicate_payments(self):
        result = self.check([(" AbC ", " νΑΙ ", "Α")], ["abc", "ABC"])
        self.assertEqual(len(result.found), 1)
        self.assertEqual(result.duplicate_payments, 1)

    def test_leading_zero_codes_remain_distinct(self):
        result = self.check([("001", "Ναι", "Α")], ["1"])
        self.assertEqual(len(result.missing), 1)
        self.assertEqual(len(result.unknown), 1)

    def test_duplicate_registry_never_marked_as_resolved(self):
        result = self.check([("001", "Ναι", "Α"), ("001", "Όχι", "Β")], ["001"])
        self.assertEqual(len(result.issues), 2)
        self.assertEqual(result.found, [])
        self.assertEqual(result.missing, [])
        self.assertEqual(result.unknown, [])

    def test_empty_codes_and_status_are_issues(self):
        result = self.check([("", "Ναι", "Α"), ("002", "", "Β")], [""])
        self.assertEqual(len(result.issues), 3)
        self.assertEqual(result.found, [])
        self.assertEqual(result.missing, [])

    def test_empty_payment_file_lists_every_active_customer(self):
        result = self.check([("001", "Ναι", "Α"), ("002", "Ναι", "Β")], [])
        self.assertEqual(len(result.missing), 2)

    def test_invalid_configuration_rejected(self):
        registry = table([("001", "Ναι", "Α")])
        payments = table([], ("Κωδικός",))
        for code, status, active in [(0, 0, {"Ναι"}), (9, 1, {"Ναι"}), (0, 1, set()), (0, 1, {"Wrong"})]:
            with self.subTest(code=code, status=status, active=active), self.assertRaises(ValueError):
                compare(registry, payments, code, status, 0, active)


class ExcelTests(unittest.TestCase):
    def test_formatted_numeric_codes_and_header_offset(self):
        with TemporaryDirectory() as folder:
            path = Path(folder) / "input.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Πελάτες"
            sheet.append(["Τίτλος"])
            sheet.append(["Κωδικός", "Συμβόλαιο"])
            sheet.append([7, "Ναι"])
            sheet["A3"].number_format = "0000"
            sheet.append([None, None])
            sheet.append(["0008", "Όχι"])
            workbook.save(path)
            workbook.close()
            self.assertEqual(sheets(str(path)), ["Πελάτες"])
            data = read_table(str(path), "Πελάτες", 2)
            self.assertEqual(data.rows[0].values[0], "0007")
            self.assertEqual(data.rows[1].values[0], "0008")
            self.assertEqual(data.rows[1].number, 5)

    def test_blank_headers_and_legacy_format_rejected(self):
        with TemporaryDirectory() as folder:
            path = Path(folder) / "empty.xlsx"
            workbook = Workbook()
            workbook.save(path)
            workbook.close()
            with self.assertRaises(ValueError):
                read_table(str(path), "Sheet")
            with self.assertRaises(ValueError):
                sheets(str(Path(folder) / "legacy.xls"))

    def test_report_preserves_text_and_does_not_create_formulas(self):
        with TemporaryDirectory() as folder:
            path = Path(folder) / "report.xlsx"
            result = Result(missing=[Entry("001", "=1+1", 2)])
            export_report(str(path), result, {"Πηγή": "=example"})
            workbook = load_workbook(path)
            try:
                self.assertEqual(workbook["Λείπουν"]["A2"].value, "001")
                self.assertEqual(workbook["Λείπουν"]["B2"].value, "=1+1")
                self.assertEqual(workbook["Λείπουν"]["B2"].data_type, "s")
                self.assertEqual(workbook["Σύνοψη"]["B2"].data_type, "s")
                self.assertEqual(len(workbook.sheetnames), 5)
            finally:
                workbook.close()

    def test_full_excel_workflow(self):
        with TemporaryDirectory() as folder:
            registry_path = Path(folder) / "registry.xlsx"
            payments_path = Path(folder) / "payments.xlsx"
            for path, rows in [
                (registry_path, [("Κωδικός", "Συμβόλαιο", "Όνομα"), ("001", "Ναι", "Α"), ("002", "Ναι", "Β")]),
                (payments_path, [("Κωδικός",), ("001",), ("999",)]),
            ]:
                workbook = Workbook()
                for row in rows:
                    workbook.active.append(row)
                workbook.save(path)
                workbook.close()
            registry = read_table(str(registry_path), "Sheet")
            payments = read_table(str(payments_path), "Sheet")
            result = compare(registry, payments, 0, 1, 0, {"Ναι"}, 2)
            report = Path(folder) / "report.xlsx"
            export_report(str(report), result, {})
            workbook = load_workbook(report, data_only=True)
            try:
                self.assertEqual(workbook["Λείπουν"]["A2"].value, "002")
                self.assertEqual(workbook["Βρέθηκαν"]["A2"].value, "001")
                self.assertEqual(workbook["Άγνωστοι κωδικοί"]["A2"].value, "999")
            finally:
                workbook.close()


if __name__ == "__main__":
    unittest.main()

